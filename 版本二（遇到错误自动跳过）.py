import os
import re
from docx import Document
from openpyxl import Workbook

# 文件夹路径
folder_path = "F:\溯源报告\报告"
# 输出 Excel 文件名
output_file_name = "ip_addresses.xlsx"

wb = Workbook()
ws = wb.active
ws.append(["File Name", "IP Address", "Error"])

for filename in os.listdir(folder_path):
    if not filename.endswith(".docx"):
        continue
    try:
        document = Document(os.path.join(folder_path, filename))
        ip_address = ""
        for paragraph in document.paragraphs:
            match = re.search(r"([0-9]{1,3}\.){3}[0-9]{1,3}", paragraph.text)
            if match:
                ip_address = match.group(0)
                break
        if ip_address:
            ws.append([filename, ip_address, ""])
    except Exception as e:
        ws.append([filename, "", str(e)])
        print(f"处理文件 {filename} 时出现错误：{str(e)}")
        continue

wb.save(output_file_name)
